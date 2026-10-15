"""Helpers for building .xlsx workbooks with openpyxl (mk-xls skill).

Covers what openpyxl leaves to the caller: a styled header + data block in
one call, column auto-fit (Korean characters count double), Excel tables
without the auto-filter clash, a workbook-wide Korean font, listing formulas
(openpyxl never calculates them), and never overwriting an existing output.
"""
from __future__ import annotations

import unicodedata
from copy import copy
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_KO_FONT = "맑은 고딕"
_THIN = Side(style="thin", color="BFBFBF")


# -- writing data ----------------------------------------------------------


def write_table(ws, headers, rows, start_row: int = 1, start_col: int = 1,
                number_formats: dict | None = None, header_fill: str = "1F4E79",
                header_font_color: str = "FFFFFF", as_table: str | None = None,
                freeze: bool = True, borders: bool = True) -> str:
    """Write a header row and data rows; return the block's range, e.g. "A1:D20".

    number_formats maps a header name to a format such as "#,##0" or "0.0%".
    as_table: give a table name (letters, digits, underscores) to turn the
    block into an Excel table; otherwise a plain auto-filter is added. Only one
    of the two is applied because a sheet auto-filter over a table makes Excel
    report the file as damaged.
    """
    number_formats = number_formats or {}
    header_style = Font(bold=True, color=header_font_color)
    fill = PatternFill("solid", start_color=header_fill)
    border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN) if borders else None

    for j, name in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + j, value=name)
        cell.font = header_style
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if border:
            cell.border = border

    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = ws.cell(row=start_row + i, column=start_col + j, value=value)
            fmt = number_formats.get(headers[j])
            if fmt:
                cell.number_format = fmt
            if border:
                cell.border = border

    last_row = start_row + max(len(rows), 1)
    ref = (f"{get_column_letter(start_col)}{start_row}:"
           f"{get_column_letter(start_col + len(headers) - 1)}{last_row}")

    if as_table:
        add_excel_table(ws, ref, as_table)
    else:
        ws.auto_filter.ref = ref
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    return ref


def add_excel_table(ws, ref: str, name: str, style: str = "TableStyleMedium2") -> Table:
    """Add an Excel table over `ref`. Header cells must be unique strings."""
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True,
                                          showColumnStripes=False)
    ws.add_table(table)
    return table


# -- layout and fonts ------------------------------------------------------


def _display_width(value) -> int:
    if value is None:
        return 0
    text = str(value)
    if text.startswith("="):
        return 10  # formula: the result is unknown, assume a typical number
    return max(
        sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)
        for line in text.splitlines() or [""]
    )


def autofit_columns(ws, min_width: float = 8, max_width: float = 60) -> None:
    """Set each column's width from its longest value (Korean counts as 2)."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), _display_width(cell.value))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))


def apply_font(wb, name: str = DEFAULT_KO_FONT) -> None:
    """Use `name` for every non-empty cell, keeping size, bold, and color.

    Excel's default font (Calibri) has no Korean glyphs, so Korean falls back
    to a system font; setting the font explicitly keeps the look consistent.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    font = copy(cell.font)
                    font.name = name
                    cell.font = font


# -- checks ----------------------------------------------------------------


def list_formulas(ws) -> list[tuple[str, str]]:
    """Return (coordinate, formula) for every formula cell in `ws`.

    openpyxl stores formulas but never calculates them, so a saved file has no
    cached results until Excel or LibreOffice opens and saves it.
    """
    return [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]


# -- output ----------------------------------------------------------------


def next_available_path(path: str | Path) -> Path:
    """Return `path`, or `name_2.ext`, `name_3.ext`, ... if it already exists."""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1
