"""Helpers for building .docx documents with python-docx (mk-doc skill).

Covers what python-docx has no public API for: Korean (East Asian) fonts,
A4 pages (the built-in template is US Letter), hyperlinks, page-number and
table-of-contents fields, cell shading and repeating table headers,
find-and-replace across runs, and never overwriting an existing output.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run

DEFAULT_KO_FONT = "맑은 고딕"
_THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


# -- document setup --------------------------------------------------------


def new_document(template: str | Path | None = None, font: str = DEFAULT_KO_FONT,
                 size: float = 11, a4: bool = True, margin_cm: float = 2.5):
    """Return a new Document with the Korean font applied to all styles.

    With a template, its page setup is kept. Without one, the built-in
    template's US Letter page is changed to A4 with `margin_cm` margins.
    """
    doc = Document(str(template)) if template is not None else Document()
    if template is None and a4:
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.left_margin = section.right_margin = Cm(margin_cm)
            section.top_margin = section.bottom_margin = Cm(margin_cm)
    set_style_fonts(doc, font)
    doc.styles["Normal"].font.size = Pt(size)
    return doc


# -- fonts -----------------------------------------------------------------


def _set_rfonts(rPr, name: str) -> None:
    rFonts = rPr.get_or_add_rFonts()
    for attr in _THEME_FONT_ATTRS:  # theme fonts win over explicit ones
        rFonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def set_run_font(run, name: str | None = None, size: float | None = None,
                 bold: bool | None = None, color: str | None = None) -> None:
    """Set a run's font, including the East Asian typeface.

    python-docx's `font.name` leaves w:eastAsia alone, so Korean text keeps
    the style or theme font unless it is set too. `color` is hex like "1F4E79".
    """
    if name is not None:
        _set_rfonts(run._element.get_or_add_rPr(), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_fonts(doc, name: str = DEFAULT_KO_FONT) -> None:
    """Use `name` in the document defaults and every paragraph/character style.

    Heading styles in the built-in template point at theme fonts, which would
    otherwise override the font set on Normal.
    """
    rPrDefault = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault"))
    if rPrDefault is not None:
        rPr = rPrDefault.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            rPrDefault.append(rPr)
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in _THEME_FONT_ATTRS:
            rFonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), name)
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            _set_rfonts(style.element.get_or_add_rPr(), name)


def _iter_paragraphs(doc):
    def from_container(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from from_container(cell)

    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from from_container(part)


def apply_korean_font(doc, name: str = DEFAULT_KO_FONT) -> None:
    """Set `name` on all styles and on every run in body, tables, headers, footers."""
    set_style_fonts(doc, name)
    for para in _iter_paragraphs(doc):
        for run in para.runs:
            set_run_font(run, name)


# -- links and fields ------------------------------------------------------


def add_hyperlink(paragraph, text: str, url: str, color: str = "0563C1",
                  underline: bool = True) -> Run:
    """Append a clickable link to `paragraph` and return its run for styling."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    link.append(r)
    paragraph._p.append(link)
    run = Run(r, paragraph)
    run.text = text
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.underline = underline
    return run


def _field_runs(paragraph, instr: str, placeholder: str = "") -> None:
    def fld_char(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            el.set(qn("w:dirty"), "true")  # ask Word to refresh on open
        return el

    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    for el in (fld_char("begin"), instr_el, fld_char("separate")):
        paragraph.add_run()._r.append(el)
    paragraph.add_run(placeholder)
    paragraph.add_run()._r.append(fld_char("end"))


def add_page_number(section, pattern: str = "{PAGE} / {NUMPAGES}",
                    align=WD_ALIGN_PARAGRAPH.CENTER, header: bool = False) -> None:
    """Put a page number in the section footer (or header).

    `pattern` mixes text and the fields {PAGE} and {NUMPAGES},
    e.g. "- {PAGE} -" or "{PAGE} / {NUMPAGES}".
    """
    part = section.header if header else section.footer
    part.is_linked_to_previous = False
    para = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    para.alignment = align
    rest = pattern
    while rest:
        start = rest.find("{")
        if start == -1:
            para.add_run(rest)
            break
        if start:
            para.add_run(rest[:start])
        end = rest.index("}", start)
        _field_runs(para, rest[start + 1:end], "1")
        rest = rest[end + 1:]


def add_toc(doc, levels: str = "1-3",
            placeholder: str = "목차: Word에서 이 영역을 마우스 오른쪽 단추로 클릭한 뒤 '필드 업데이트'를 선택하세요."):
    """Insert a table-of-contents field built from Heading 1..n.

    python-docx cannot compute page numbers, so the TOC shows `placeholder`
    until Word updates fields. The document is also flagged so Word offers to
    update fields when it opens.
    """
    para = doc.add_paragraph()
    _field_runs(para, f'TOC \\o "{levels}" \\h \\z \\u', placeholder)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        flag = OxmlElement("w:updateFields")
        flag.set(qn("w:val"), "true")
        settings.append(flag)
    return para


# -- tables ----------------------------------------------------------------


def set_cell_background(cell, color: str) -> None:
    """Fill a table cell with a hex color such as "1F4E79"."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)


def add_table(doc, headers, rows, style: str = "Table Grid", header_fill: str = "1F4E79",
              header_font_color: str = "FFFFFF", col_widths_cm=None):
    """Add a table with a shaded header row that repeats on every page."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        set_cell_background(cell, header_fill)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(header_font_color)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = "" if value is None else str(value)
    if col_widths_cm:
        for row in table.rows:
            for cell, width in zip(row.cells, col_widths_cm):
                cell.width = Cm(width)
    return table


# -- editing ---------------------------------------------------------------


def _replace_in_paragraph(paragraph, old: str, new: str) -> int:
    count, search_from = 0, 0
    while True:
        runs = paragraph.runs
        texts = [r.text for r in runs]
        full = "".join(texts)
        idx = full.find(old, search_from)
        if idx == -1:
            return count
        starts, pos = [], 0
        for t in texts:
            starts.append(pos)
            pos += len(t)
        end = idx + len(old)

        def run_at(char_index):  # the run that holds this character
            return next(i for i, s in enumerate(starts) if s <= char_index < s + len(texts[i]))

        si, ei = run_at(idx), run_at(end - 1)
        s_off, e_off = idx - starts[si], end - starts[ei]
        if si == ei:
            runs[si].text = texts[si][:s_off] + new + texts[si][e_off:]
        else:
            runs[si].text = texts[si][:s_off] + new
            for k in range(si + 1, ei):
                runs[k].text = ""
            runs[ei].text = texts[ei][e_off:]
        count += 1
        search_from = idx + len(new)


def replace_text(doc, old: str, new: str) -> int:
    """Replace `old` with `new` everywhere; return the number of replacements.

    Works even when Word split the text across runs; the replacement takes the
    formatting of the run where the match starts. Text inside hyperlinks and
    text boxes is not searched.
    """
    if not old:
        raise ValueError("old must not be empty")
    return sum(_replace_in_paragraph(p, old, new) for p in _iter_paragraphs(doc))


# -- output ----------------------------------------------------------------


def next_available_path(path: str | Path) -> Path:
    """Return `path`, or `name_2.ext`, `name_3.ext`, ... if it already exists."""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1
