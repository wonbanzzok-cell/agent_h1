"""Helpers for building .pptx decks with python-pptx (mk-ppt skill).

Covers what python-pptx has no public API for: Korean (East Asian) fonts,
deleting / duplicating / reordering slides, 16:9 decks whose layouts actually
fill the wider slide, and never overwriting an existing output file.
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.util import Emu, Pt

DEFAULT_KO_FONT = "Noto Sans KR"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# -- deck setup ------------------------------------------------------------


def new_presentation(template: str | Path | None = None, widescreen: bool = True):
    """Return a new Presentation.

    With a template, its slide size and layouts are kept as they are. Without
    one, python-pptx's built-in 4:3 template is widened to 16:9 and its master
    and layouts are stretched to match, so placeholders span the full width.
    """
    if template is not None:
        return Presentation(str(template))
    prs = Presentation()
    if widescreen:
        old_width = prs.slide_width
        prs.slide_width = Emu(12192000)  # 13.333 in
        prs.slide_height = Emu(6858000)  # 7.5 in
        _stretch_x(prs, prs.slide_width / old_width)
    return prs


def _stretch_x(prs, factor: float) -> None:
    """Scale x offsets and widths on the slide master and every layout."""
    parts = [prs.slide_master] + list(prs.slide_layouts)
    for part in parts:
        tree = part.shapes._spTree
        for off in tree.iter(qn("a:off")):
            off.set("x", str(int(int(off.get("x")) * factor)))
        for ext in tree.iter(qn("a:ext")):
            if ext.get("cx") is not None:
                ext.set("cx", str(int(int(ext.get("cx")) * factor)))


def add_title_slide(prs, title: str, subtitle: str | None = None):
    """Add a slide from layout 0 (Title Slide)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle is not None and 1 in [p.placeholder_format.idx for p in slide.placeholders]:
        slide.placeholders[1].text = subtitle
    return slide


def add_bullet_slide(prs, title: str, items):
    """Add a slide from layout 1 (Title and Content) with bullet items.

    Each item is either "text" (level 0) or ("text", level) with level 0-8.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = text
        para.level = level
    return slide


# -- fonts -----------------------------------------------------------------


def set_font(font, name: str | None = None, size: float | None = None,
             bold: bool | None = None, color: str | None = None) -> None:
    """Set a run's or paragraph's font, including the East Asian typeface.

    `font` is `run.font` or `paragraph.font`. python-pptx's `font.name` only
    writes <a:latin>, so Korean text keeps the theme font unless <a:ea> is set
    too. `size` is in points, `color` is a hex string such as "1F4E79".
    """
    if name is not None:
        font.name = name
        rPr = font._rPr
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = rPr.makeelement(qn("a:ea"), {})
            rPr.find(qn("a:latin")).addnext(ea)
        ea.set("typeface", name)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)


def _iter_text_frames(shapes):
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _iter_text_frames(shape.shapes)
        if shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def apply_korean_font(prs, name: str = DEFAULT_KO_FONT) -> None:
    """Apply `name` as Latin and East Asian font to every run in the deck.

    Charts are not touched; style chart fonts through the chart API.
    """
    for slide in prs.slides:
        for tf in _iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                for run in para.runs:
                    set_font(run.font, name)


# -- pink theme ------------------------------------------------------------

PINK_THEME = {
    "background": "FFF0F5",  # lavender blush
    "primary": "FF69B4",     # hot pink: titles
    "secondary": "F4A6C1",   # rose: accent bars, shapes
    "accent": "C2185B",      # deep pink: highlights
    "text": "4A4A4A",        # dark gray: body text
}


def set_slide_background(slide, color: str) -> None:
    """Fill `slide`'s background with a solid hex color such as "FFF0F5"."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


def add_accent_bar(slide, prs, color: str = PINK_THEME["secondary"],
                   top: Emu = Emu(1600200), height: Emu = Emu(73152)):
    """Add a thin full-width bar (default: just under the title) and return it."""
    margin = Emu(457200)  # 0.5 in
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, margin, top,
                                 prs.slide_width - 2 * margin, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(color)
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def apply_pink_theme(prs, theme: dict = PINK_THEME,
                     font: str = DEFAULT_KO_FONT) -> None:
    """Apply the pink theme to every slide: background, title and body colors, font.

    Titles get `theme["primary"]` in bold, other text gets `theme["text"]`.
    Call it after all content is added. Charts are not touched.
    """
    for slide in prs.slides:
        set_slide_background(slide, theme["background"])
        title = slide.shapes.title
        for shape in slide.shapes:
            is_title = title is not None and shape.shape_id == title.shape_id
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if is_title:
                        set_font(run.font, font, bold=True, color=theme["primary"])
                    else:
                        set_font(run.font, font, color=theme["text"])
        for tf in _iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                for run in para.runs:
                    if run.font._rPr.find(qn("a:ea")) is None:
                        set_font(run.font, font)


# -- slide management ------------------------------------------------------


def _slide_index(prs, slide) -> int:
    return [s.slide_id for s in prs.slides].index(slide.slide_id)


def delete_slide(prs, slide) -> None:
    """Remove `slide` from the deck."""
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst.sldId_lst[_slide_index(prs, slide)]
    rId = sldId.rId
    sldIdLst.remove(sldId)
    prs.part.drop_rel(rId)


def move_slide(prs, slide, new_index: int) -> None:
    """Move `slide` to position `new_index` (0-based)."""
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst.sldId_lst[_slide_index(prs, slide)]
    sldIdLst.remove(sldId)
    sldIdLst.insert(new_index, sldId)


def duplicate_slide(prs, slide):
    """Append a copy of `slide` (same layout, shapes, background, images, links).

    Speaker notes are not copied. A chart in the copy shares its chart part
    with the original, so build charts on each slide instead of duplicating.
    """
    new = prs.slides.add_slide(slide.slide_layout)
    new_tree = new.shapes._spTree
    for shape in list(new.shapes):
        new_tree.remove(shape._element)

    rid_map = {}
    for rId, rel in slide.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)

    skip = {qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")}
    for el in slide.shapes._spTree.iterchildren():
        if el.tag not in skip:
            new_tree.append(copy.deepcopy(el))

    bg = slide._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))

    for el in new._element.iter():
        for attr, value in el.attrib.items():
            if attr.startswith("{%s}" % _R_NS) and value in rid_map:
                el.set(attr, rid_map[value])
    return new


# -- output ----------------------------------------------------------------


def next_available_path(path: str | Path) -> Path:
    """Return `path`, or `name_2.ext`, `name_3.ext`, ... if it already exists."""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1
