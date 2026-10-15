"""Build slides/클로드_회사_소개.pptx. Run: python scripts/build_anthropic_company_pptx.py"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "mk-ppt" / "scripts"))

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt, Inches
from pptx_helpers import new_presentation, set_font, apply_korean_font, next_available_path

OUT = next_available_path(ROOT / "slides" / "클로드_회사_소개.pptx")

ORANGE = "E8742A"
DARK = "333333"
BEIGE = "FBF1E8"
WHITE = "FFFFFF"
BLANK = 6  # layout index of the blank layout


def rgb(hex_):
    return RGBColor.from_string(hex_)


def add_box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    box = slide.shapes.add_shape(shape, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = rgb(fill)
    box.line.fill.background()
    return box


def add_text(slide, x, y, w, h, lines, size=18, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """Text box; `lines` is a list of str or (str, size, color, bold) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        text, sz, col, b = (line, size, color, bold) if isinstance(line, str) else line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        set_font(run.font, size=sz, bold=b, color=col)
    return tb


def content_slide(prs, title, number):
    """Blank slide with an orange top bar, title, and page number."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    W = prs.slide_width
    add_box(slide, 0, 0, W, Inches(0.25), ORANGE)
    add_box(slide, Inches(0.6), Inches(0.6), Inches(0.12), Inches(0.6), ORANGE)
    add_text(slide, Inches(0.85), Inches(0.5), Inches(11), Inches(0.8),
             [title], size=30, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, W - Inches(1.2), Inches(6.9), Inches(0.8), Inches(0.4),
             [str(number)], size=12, color=ORANGE, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    W, H = prs.slide_width, prs.slide_height
    add_box(slide, 0, 0, W, H, ORANGE)
    add_box(slide, Inches(0.9), Inches(3.55), Inches(1.5), Inches(0.08), WHITE)
    add_text(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.5),
             [("Anthropic", 54, WHITE, True), ("Claude를 만드는 회사", 32, WHITE, True)],
             space_after=0)
    add_text(slide, Inches(0.9), Inches(3.9), Inches(11.5), Inches(0.8),
             ["안전하고 유익한 AI를 위한 연구 기업"], size=22, color=WHITE)
    add_text(slide, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.5),
             ["회사 소개 발표자료"], size=14, color=BEIGE)


def slide_overview(prs):
    slide = content_slide(prs, "회사 개요", 2)
    rows = [
        ("설립", "2021년"),
        ("본사", "미국 캘리포니아주 샌프란시스코"),
        ("공동 창업", "Dario Amodei(CEO), Daniela Amodei(President) 등\nOpenAI 출신 연구진"),
        ("법인 형태", "공익법인(Public Benefit Corporation)"),
        ("주요 사업", "AI 안전 연구 및 대규모 언어 모델 Claude 개발·서비스"),
    ]
    y = Inches(1.6)
    for label, value in rows:
        h = Inches(0.95)
        add_box(slide, Inches(0.85), y, Inches(2.4), h - Inches(0.12), ORANGE)
        add_text(slide, Inches(0.85), y, Inches(2.4), h - Inches(0.12), [label],
                 size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        add_box(slide, Inches(3.35), y, Inches(9.1), h - Inches(0.12), BEIGE)
        add_text(slide, Inches(3.6), y, Inches(8.8), h - Inches(0.12), [value],
                 size=18, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        y += h


def slide_mission(prs):
    slide = content_slide(prs, "미션과 핵심 가치", 3)
    add_text(slide, Inches(0.85), Inches(1.45), Inches(11.6), Inches(0.6),
             ["“AI 전환기를 인류가 안전하게 지나도록 돕는다”"],
             size=20, color=ORANGE, bold=True)
    cards = [
        ("Constitutional AI", "명시된 원칙(헌법)을 바탕으로\n모델이 스스로 답변을 점검하고\n개선하도록 학습시키는 방법"),
        ("책임 있는 확장 정책", "Responsible Scaling Policy\n모델 능력이 커질수록\n더 강한 안전 조치를 단계적으로 적용"),
        ("해석 가능성 연구", "Interpretability\n모델 내부가 어떻게 작동하는지\n들여다보고 이해하려는 연구"),
    ]
    w, gap, x = Inches(3.7), Inches(0.25), Inches(0.85)
    for head, body in cards:
        add_box(slide, x, Inches(2.3), w, Inches(4.3), BEIGE, MSO_SHAPE.ROUNDED_RECTANGLE)
        add_box(slide, x, Inches(2.3), w, Inches(1.0), ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, x, Inches(2.3), w, Inches(1.0), [head], size=20, color=WHITE,
                 bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        add_text(slide, x + Inches(0.2), Inches(3.5), w - Inches(0.4), Inches(2.9),
                 body.split("\n"), size=16, align=PP_ALIGN.CENTER, space_after=4)
        x += w + gap


def slide_product(prs):
    slide = content_slide(prs, "주요 제품: Claude", 4)
    data = [
        ("구분", "내용"),
        ("모델 라인업", "Opus(최고 성능) · Sonnet(성능과 속도 균형) · Haiku(빠르고 경제적)"),
        ("이용 채널", "claude.ai 웹·데스크톱·모바일 앱, Claude API(개발자용)"),
        ("개발 도구", "Claude Code: 터미널·IDE에서 쓰는 코딩 에이전트"),
        ("주요 활용", "글쓰기·요약, 코딩, 데이터 분석, 문서 작성, 업무 자동화"),
    ]
    tbl = slide.shapes.add_table(len(data), 2, Inches(0.85), Inches(1.6),
                                 Inches(11.6), Inches(4.6)).table
    tbl.columns[0].width = Inches(2.6)
    tbl.columns[1].width = Inches(9.0)
    for r, (a, b) in enumerate(data):
        for c, text in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.text = text
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(ORANGE)
                color, bold = WHITE, True
            else:
                cell.fill.fore_color.rgb = rgb(BEIGE if r % 2 else WHITE)
                color, bold = (ORANGE, True) if c == 0 else (DARK, False)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_font(run.font, size=18, bold=bold, color=color)


def slide_summary(prs):
    slide = content_slide(prs, "정리 및 전망", 5)
    points = [
        ("01", "안전을 최우선으로 하는 AI 연구 기업"),
        ("02", "Claude 모델군으로 개인·기업·개발자에게 AI 제공"),
        ("03", "안전 연구와 제품 개발을 함께 추진하며 업계 기준 제시"),
    ]
    y = Inches(1.7)
    for num, text in points:
        add_box(slide, Inches(0.85), y, Inches(0.9), Inches(0.9), ORANGE, MSO_SHAPE.OVAL)
        add_text(slide, Inches(0.85), y, Inches(0.9), Inches(0.9), [num], size=20,
                 color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        add_text(slide, Inches(2.0), y, Inches(10.4), Inches(0.9), [text], size=22,
                 anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        y += Inches(1.15)
    add_box(slide, Inches(0.85), Inches(5.4), Inches(11.6), Inches(1.1), ORANGE,
            MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, Inches(0.85), Inches(5.4), Inches(11.6), Inches(1.1),
             ["안전하고 신뢰할 수 있는 AI, Anthropic이 그리는 미래입니다"],
             size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, space_after=0)


def main():
    prs = new_presentation()  # 16:9
    slide_cover(prs)
    slide_overview(prs)
    slide_mission(prs)
    slide_product(prs)
    slide_summary(prs)
    apply_korean_font(prs)  # last, so every text run gets the font
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
